from __future__ import annotations

import base64
import io
import math
import shutil
import struct
import subprocess
import sys
import wave
import zipfile
from email.message import EmailMessage
from pathlib import Path

from docx import Document
from extract_msg import OleWriter
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches


def main() -> None:
    output = Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures/generated")
    output.mkdir(parents=True, exist_ok=True)

    create_image(output / "image.png")
    create_wav(output / "audio.wav")
    create_mp3(output / "audio.mp3")
    create_markup_files(output)
    create_docx(output / "document.docx")
    create_rtf(output / "document.rtf")
    create_xlsx(output / "table.xlsx")
    create_pptx(output / "presentation.pptx")
    (output / "readme.txt").write_text("Тестовый текстовый файл.\nВторая строка.\n", encoding="utf-8")
    (output / "data.csv").write_text("Название;Количество;Цена\nКофе;4;120.50\nЧай;7;85\n", encoding="utf-8")

    create_archives(output)
    create_email(output / "message.eml", output / "document.docx")
    create_outlook_message(output / "message.msg", output / "document.docx")
    (output / "broken.zip").write_bytes(b"PK\x03\x04this is deliberately damaged")
    convert_legacy_and_pdf(output)
    (output / ".fixtures-version").write_text("4\n", encoding="ascii")
    print(f"Generated fixtures in {output}")


def create_image(path: Path) -> None:
    image = Image.new("RGB", (1200, 700), "#15211b")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((90, 90, 1110, 610), radius=36, fill="#23372d", outline="#65e6a1", width=5)
    draw.text((150, 300), "FILE VIEWER / IMAGE", fill="#dfffea", font=None)
    image.save(path)


def create_wav(path: Path) -> None:
    sample_rate = 16_000
    frames = bytearray()
    for index in range(sample_rate):
        sample = int(8_000 * math.sin(2 * math.pi * 440 * index / sample_rate))
        frames.extend(struct.pack("<h", sample))
    with wave.open(str(path), "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(sample_rate)
        audio.writeframes(frames)


def create_mp3(path: Path) -> None:
    # A short 440 Hz tone, generated once with ffmpeg and embedded so fixture
    # creation remains fully offline and does not add an ffmpeg runtime dependency.
    encoded = """
    SUQzBAAAAAAAI1RTU0UAAAAPAAADTGF2ZjYyLjEzLjEwMgAAAAAAAAAAAAAA//MwxAAN2EqkX0kYAgrlttAARo0aNG3sIIxWCYBwNhsnbmcEAYlwfB8H+CDsHz/DEpygf4OHMQA/rBw5iAH9YEOaA/0KLzAIDaT9ydEx//MyxAkRGNp8AZxoAIjAwU+2sGUBcUBQxVNjToMMFgwStSoDJIQckHSCugrP44hhhLvyRHqPUy/y8XjEul35YGhKEv4NCU6DVcaFrrskgP/zMMQGDvBaLAHeAAAEYEoFhgXiiGDaHkYfoLBhcjyGWGSsf8K7JlwEWkQ5hgKgGGBqB8EAXmAQAGXeaWyXn////6knUSYD/3q69Rjrxv/zMsQLDmhOIFTn9ECerDkFjAGgBE12IdwCYo8KaDCqK3mwgqzlplD9lnd92uR/w2KRc3Z7a8/1CjoESoABHFmfgV0CpgOiIOk/rETAFAD/8zDEEwvISjG0P7IkD3DTxE3lmwNO3+MpFk+ojf+ujzWqu//9n7OP9dU0SA/W9VWxGL0Qo2rDFxjAIgCM2j4qWAeUWJL9eait5s3qrL3/8zDEJA14Thws7/RAlFClNTq/+T/4E39XTdZ1NYkcJBFalAMmA1hDSGSpS2pgEYBQbP+XInzACxFfzyyKwESEXYvv3+vkaunnlnvsq/3/8zLELwzYShxCV/QkLEer8lUccfr8rr1GZOinKsMXiMAaAJTWkCvgBvRYMzWBp23tj1UfM3bfv/8l/vfZ/17f4sqAma4/9byeEzbykvlh//MwxD0L2E4gKuf0QEvCYA2AUGsll4p3QAsFZtAs9eFcXRpUvZ+un+iu7/R0/tdrr7r9b1VbsZRSmXBSGMABAADAKAEM3I5BIOADQwaV//MwxE4McEog6uf0QIXWkNsWd+37r1VPP/913+tn7fZHVewE2eHf7b/gDLMK7/I6FXeYNuU2vWbq09M33f+Q/ravX+6qpn+NHD/3qq8R//MyxF0MOEoYAO/2QIE8D7oSUxiyRgDwCCbCYf3HoFhw5ar9Tt9Uqnr/7o+//WvxrP7hi2+3X2/+v/YqJUMP/LdZ4TU/sQDJigAABMAfAf/zMMRuCYA2in4LdiIQ199C1POMDBK1n5l14VpHchrt6/R6fqf/26f2PxtF6qe3DDSCziHEwOAuzAABfMLMJsw/BSDEgBzOehugxIQEDP/zMsSJDQj6HCj+hGQyAsDBeA2MC4AgwCQE1a2yqxxix/////fViG2VhAQLMHCkaGNNdAoHNWtwxeeDJiRj0qOEChLgz0STEQVt4AcEF9D/8zDElgzAShxK5/RA8UQbkhjHiCBVNw/ULhRkOQQ6bsLNGZJQgPk2bl8vuSI5RMnCLfL5vL7snJo2Mi8mZf1NQZBCgbHTJzFIS/u8qNP/8zDEpA5IWjQBXgAApIVUTEFNRTMuMTAxIChiZXRhIDMpqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqr/8zLEqxs5jnApnJAAqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqq//MwxIAAAANIAcAAAKqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqqq
    """
    path.write_bytes(base64.b64decode(encoded))


def create_markup_files(output: Path) -> None:
    (output / "document.xml").write_text(
        """<?xml version="1.0" encoding="UTF-8"?>
<catalog>
  <document id="42">
    <title>Тестовый XML</title>
    <status>ready</status>
  </document>
</catalog>
""",
        encoding="utf-8",
    )
    html = """<!doctype html>
<html lang="ru">
  <head><title>Тестовый HTML</title></head>
  <body>
    <h1>Безопасный HTML</h1>
    <p onclick="alert('blocked')">Скрипты и внешние ресурсы удаляются.</p>
    <img src="http://tracking.invalid/pixel.gif" alt="blocked tracker">
    <script>alert('blocked')</script>
  </body>
</html>
"""
    (output / "page.html").write_text(html, encoding="utf-8")
    (output / "page.htm").write_text(html, encoding="utf-8")


def create_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Тестовый документ", 0)
    document.add_paragraph("Этот DOCX создан автоматически для локальной проверки предпросмотра.")
    document.add_heading("Проверяем", level=1)
    for value in ("кириллицу и шрифты", "списки", "таблицы", "конвертацию в PDF"):
        document.add_paragraph(value, style="List Bullet")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Поле"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "Статус"
    table.cell(1, 1).text = "Готово"
    table.cell(2, 0).text = "Версия"
    table.cell(2, 1).text = "0.1"
    document.save(path)


def create_rtf(path: Path) -> None:
    content = (
        r"{\rtf1\ansi\ansicpg1251\deff0"
        r"{\fonttbl{\f0\fswiss Liberation Sans;}}"
        r"\viewkind4\uc1\pard\f0\fs34\b RTF preview\b0\par"
        r"\fs24 This file is converted to PDF by LibreOffice.\par}"
    )
    path.write_text(content, encoding="ascii")


def create_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Товары"
    sheet.append(["Название", "Количество", "Цена", "Сумма"])
    for row, values in enumerate((("Кофе", 4, 120.5), ("Чай", 7, 85), ("Какао", 3, 150)), start=2):
        sheet.append([*values, f"=B{row}*C{row}"])
    summary = workbook.create_sheet("Описание")
    summary.append(["Книга", "Локальная тестовая таблица"])
    summary.append(["Листов", 2])
    workbook.save(path)


def create_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "File Viewer"
    slide.placeholders[1].text = "Тестовая презентация"
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Вторая страница"
    box = second.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    box.text_frame.text = "PPTX преобразуется в PDF через LibreOffice."
    presentation.save(path)


def create_archives(output: Path) -> None:
    nested_buffer = io.BytesIO()
    with zipfile.ZipFile(nested_buffer, "w", zipfile.ZIP_DEFLATED) as nested:
        nested.writestr("inside/note.txt", "Вложенный архив открыт успешно.\n")
    archive_path = output / "archive.zip"
    with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("docs/readme.txt", "Файл внутри архива.\n")
        archive.write(output / "table.xlsx", "docs/table.xlsx")
        archive.write(output / "audio.mp3", "media/audio.mp3")
        archive.writestr("nested.zip", nested_buffer.getvalue())
        archive.writestr("empty/", b"")
    shutil.copy2(archive_path, output / "archive.custom-package")
    try:
        subprocess.run(
            ["7z", "a", "-t7z", "-pviewer", "-mhe=on", str(output / "protected.7z"), str(output / "readme.txt")],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
    except (FileNotFoundError, subprocess.CalledProcessError):
        pass


def create_email(path: Path, attachment: Path) -> None:
    message = EmailMessage()
    message["From"] = "sender@example.internal"
    message["To"] = "receiver@example.internal"
    message["Subject"] = "Проверка EML и вложения"
    message.set_content("Текстовая версия тестового письма.")
    message.add_alternative(
        """
        <html><body>
          <h2>Тестовое письмо</h2>
          <p>Внешнее изображение и скрипт должны быть заблокированы.</p>
          <img src="http://tracking.invalid/pixel.gif">
          <script>alert('not allowed')</script>
        </body></html>
        """,
        subtype="html",
    )
    message.add_attachment(
        attachment.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename="Вложенный документ.docx",
    )
    path.write_bytes(message.as_bytes())


def create_outlook_message(path: Path, attachment: Path) -> None:
    writer = OleWriter()
    transport_headers = (
        "From: Outlook Sender <sender@example.internal>\r\n"
        "To: receiver@example.internal\r\n"
        "Subject: Проверка Outlook MSG и вложения\r\n"
        "MIME-Version: 1.0\r\n"
        "\r\n"
    )
    root_values = {
        "001A001F": _msg_unicode("IPM.Note"),
        "0037001F": _msg_unicode("Проверка Outlook MSG и вложения"),
        "007D001F": _msg_unicode(transport_headers),
        "1000001F": _msg_unicode("Текстовая версия тестового Outlook MSG."),
        "10130102": "<html><body><h2>Outlook MSG</h2><p>MSG открыт без Outlook.</p></body></html>".encode("utf-8"),
        "0C1A001F": _msg_unicode("Outlook Sender"),
        "0C1F001F": _msg_unicode("sender@example.internal"),
        "5D01001F": _msg_unicode("sender@example.internal"),
        "0E04001F": _msg_unicode("receiver@example.internal"),
    }
    root_properties = b"".join(
        _msg_variable_property(name, len(data)) for name, data in root_values.items()
    )
    root_header = struct.pack("<QIIIIQ", 0, 0, 1, 0, 1, 0)
    writer.addEntry("__properties_version1.0", root_header + root_properties)
    for name, data in root_values.items():
        writer.addEntry(f"__substg1.0_{name}", data)
    for name in ("00020102", "00030102", "00040102"):
        writer.addEntry(["__nameid_version1.0", f"__substg1.0_{name}"], b"")

    attachment_values = {
        "3707001F": _msg_unicode("Вложенный документ.docx"),
        "370E001F": _msg_unicode(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        "37010102": attachment.read_bytes(),
    }
    attachment_properties = struct.pack("<HHII", 0x0003, 0x3705, 0, 1) + b"\x00" * 4
    attachment_properties += b"".join(
        _msg_variable_property(name, len(data)) for name, data in attachment_values.items()
    )
    storage = "__attach_version1.0_#00000000"
    writer.addEntry([storage, "__properties_version1.0"], b"\x00" * 8 + attachment_properties)
    for name, data in attachment_values.items():
        writer.addEntry([storage, f"__substg1.0_{name}"], data)
    writer.write(path)


def _msg_unicode(value: str) -> bytes:
    # extract-msg exposes PT_UNICODE streams verbatim. Omitting the optional
    # terminator keeps the generated fixture headers free of trailing NULs.
    return value.encode("utf-16-le")


def _msg_variable_property(name: str, size: int) -> bytes:
    property_id = int(name[:4], 16)
    property_type = int(name[4:], 16)
    return struct.pack("<HHIII", property_type, property_id, 0, size, 0)


def convert_legacy_and_pdf(output: Path) -> None:
    conversions = [
        (output / "document.docx", "pdf"),
        (output / "document.docx", "doc"),
        (output / "table.xlsx", "xls"),
        (output / "presentation.pptx", "ppt"),
    ]
    for source, target in conversions:
        try:
            subprocess.run(
                ["soffice", "--headless", "--convert-to", target, "--outdir", str(output), str(source)],
                check=True,
                timeout=90,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
        except (FileNotFoundError, subprocess.CalledProcessError, subprocess.TimeoutExpired):
            continue


if __name__ == "__main__":
    main()
