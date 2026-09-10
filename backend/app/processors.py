from __future__ import annotations

import asyncio
import codecs
import csv
import re
import shutil
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from tempfile import mkdtemp

import bleach
import rarfile
from fastapi import HTTPException
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError

from .config import Settings
from .container_formats import list_7z, list_email, list_rar, list_zip, write_json
from .detection import DetectedFormat


@dataclass(frozen=True, slots=True)
class ProcessResult:
    renderer: str
    detected_type: str
    mime_type: str
    asset_name: str | None


async def process_materialized_file(
    source_path: Path,
    detected: DetectedFormat,
    entry_dir: Path,
    settings: Settings,
) -> ProcessResult:
    entry_dir.mkdir(parents=True, exist_ok=True)
    if detected.kind == "pdf":
        target = entry_dir / "preview.pdf"
        await asyncio.to_thread(shutil.copy2, source_path, target)
        return ProcessResult("pdf", detected.label, "application/pdf", target.name)
    if detected.kind == "image":
        target = entry_dir / f"preview{detected.extension}"
        await asyncio.to_thread(shutil.copy2, source_path, target)
        return ProcessResult("image", detected.label, detected.mime_type, target.name)
    if detected.kind == "audio":
        target = entry_dir / f"preview{detected.extension}"
        await asyncio.to_thread(shutil.copy2, source_path, target)
        return ProcessResult("audio", detected.label, detected.mime_type, target.name)
    if detected.kind == "image_convert":
        return await _convert_image(source_path, entry_dir, detected)
    if detected.kind == "office":
        return await _convert_office(source_path, entry_dir, detected, settings)
    if detected.kind == "spreadsheet":
        return await _render_spreadsheet(source_path, entry_dir, detected, settings)
    if detected.kind == "delimited":
        return await _render_delimited(source_path, entry_dir, detected, settings)
    if detected.kind == "archive":
        return await _render_archive(source_path, entry_dir, detected, settings)
    if detected.kind in {"email", "outlook_email"}:
        return await _render_email(source_path, entry_dir, detected, settings)
    if detected.kind == "html":
        return await _render_html(source_path, entry_dir, detected, settings)
    if detected.kind == "xml":
        return await _render_xml(source_path, entry_dir, detected, settings)
    if detected.kind == "json":
        return await _render_json(source_path, entry_dir, detected, settings)
    if detected.kind == "text":
        return await _render_text(source_path, entry_dir, detected, settings)
    return ProcessResult("unsupported", detected.label, detected.mime_type, None)


async def _convert_office(
    source_path: Path, entry_dir: Path, detected: DetectedFormat, settings: Settings
) -> ProcessResult:
    if source_path.stat().st_size > settings.max_office_bytes:
        raise HTTPException(status_code=413, detail="Office-файл превышает лимит")
    typed_input = source_path.parent / f"input{detected.extension}"
    await asyncio.to_thread(shutil.copy2, source_path, typed_input)
    profile = Path(mkdtemp(prefix="lo-profile-", dir=source_path.parent))
    command = [
        "soffice",
        "--headless",
        "--nologo",
        "--nodefault",
        "--nolockcheck",
        f"-env:UserInstallation={profile.as_uri()}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(entry_dir),
        str(typed_input),
    ]
    await _run(command, settings.office_timeout_seconds, "Office-файл не удалось преобразовать")
    candidates = sorted(entry_dir.glob("*.pdf"))
    if not candidates:
        raise HTTPException(status_code=422, detail="LibreOffice не создал PDF")
    output = entry_dir / "preview.pdf"
    if candidates[0] != output:
        candidates[0].replace(output)
    return ProcessResult("pdf", detected.label, "application/pdf", output.name)


async def _render_spreadsheet(
    source_path: Path, entry_dir: Path, detected: DetectedFormat, settings: Settings
) -> ProcessResult:
    if source_path.stat().st_size > settings.max_office_bytes:
        raise HTTPException(status_code=413, detail="Таблица превышает лимит")
    workbook_path = source_path
    if detected.extension not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        typed_input = source_path.parent / f"input{detected.extension}"
        await asyncio.to_thread(shutil.copy2, source_path, typed_input)
        output_dir = source_path.parent / "converted"
        output_dir.mkdir(exist_ok=True)
        profile = Path(mkdtemp(prefix="lo-profile-", dir=source_path.parent))
        await _run(
            [
                "soffice",
                "--headless",
                "--nologo",
                "--nodefault",
                f"-env:UserInstallation={profile.as_uri()}",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(output_dir),
                str(typed_input),
            ],
            settings.office_timeout_seconds,
            "Таблицу не удалось преобразовать",
        )
        candidates = sorted(output_dir.glob("*.xlsx"))
        if not candidates:
            raise HTTPException(status_code=422, detail="LibreOffice не создал XLSX")
        workbook_path = candidates[0]

    data = await asyncio.to_thread(_workbook_to_dict, workbook_path, settings)
    output = entry_dir / "workbook.json"
    await asyncio.to_thread(write_json, output, data)
    return ProcessResult("spreadsheet", detected.label, "application/json", output.name)


def _workbook_to_dict(path: Path, settings: Settings) -> dict:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Таблицу не удалось прочитать") from exc
    sheets: list[dict] = []
    truncated = False
    try:
        for worksheet in workbook.worksheets:
            rows: list[list[object]] = []
            for row_index, row in enumerate(
                worksheet.iter_rows(max_col=settings.max_sheet_columns, values_only=True), start=1
            ):
                if row_index > settings.max_sheet_rows:
                    truncated = True
                    break
                values = [_json_cell(value) for value in row]
                while values and values[-1] is None:
                    values.pop()
                rows.append(values)
            sheets.append({"name": worksheet.title, "rows": rows})
    finally:
        workbook.close()
    return {"sheets": sheets, "truncated": truncated}


def _json_cell(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


async def _render_delimited(
    source_path: Path, entry_dir: Path, detected: DetectedFormat, settings: Settings
) -> ProcessResult:
    text = await asyncio.to_thread(_decode_text, source_path.read_bytes())
    delimiter = "\t" if detected.extension == ".tsv" else ","
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass
    rows: list[list[str]] = []
    truncated = False
    for index, row in enumerate(csv.reader(text.splitlines(), delimiter=delimiter), start=1):
        if index > settings.max_sheet_rows:
            truncated = True
            break
        rows.append(row[: settings.max_sheet_columns])
    output = entry_dir / "workbook.json"
    await asyncio.to_thread(
        write_json,
        output,
        {"sheets": [{"name": "Data", "rows": rows}], "truncated": truncated},
    )
    return ProcessResult("spreadsheet", detected.label, "application/json", output.name)


async def _render_archive(
    source_path: Path, entry_dir: Path, detected: DetectedFormat, settings: Settings
) -> ProcessResult:
    if source_path.stat().st_size > settings.max_archive_bytes:
        raise HTTPException(status_code=413, detail="Архив превышает лимит")
    if zipfile_is_zip(source_path):
        data = await asyncio.to_thread(list_zip, source_path, settings)
    elif rarfile.is_rarfile(source_path):
        data = await asyncio.to_thread(list_rar, source_path, settings)
    else:
        data = await list_7z(source_path, settings)
    output = entry_dir / "archive.json"
    await asyncio.to_thread(write_json, output, data)
    return ProcessResult("archive", detected.label, "application/json", output.name)


async def _render_email(
    source_path: Path, entry_dir: Path, detected: DetectedFormat, settings: Settings
) -> ProcessResult:
    if source_path.stat().st_size > settings.max_office_bytes:
        raise HTTPException(status_code=413, detail="Почтовый файл превышает лимит")
    data, _ = await asyncio.to_thread(list_email, source_path)
    data["html"] = _sanitize_email_html(data["html"])
    output = entry_dir / "email.json"
    await asyncio.to_thread(write_json, output, data)
    return ProcessResult("email", detected.label, "application/json", output.name)


async def _render_text(
    source_path: Path, entry_dir: Path, detected: DetectedFormat, settings: Settings
) -> ProcessResult:
    if source_path.stat().st_size > settings.max_text_bytes:
        raise HTTPException(status_code=413, detail="Текстовый файл превышает лимит")
    content = await asyncio.to_thread(_decode_text, source_path.read_bytes())
    output = entry_dir / "preview.txt"
    await asyncio.to_thread(output.write_text, content, encoding="utf-8")
    return ProcessResult("text", detected.label, "text/plain; charset=utf-8", output.name)


async def _render_xml(
    source_path: Path, entry_dir: Path, detected: DetectedFormat, settings: Settings
) -> ProcessResult:
    if source_path.stat().st_size > settings.max_text_bytes:
        raise HTTPException(status_code=413, detail="XML-файл превышает лимит")
    content = await asyncio.to_thread(_decode_text, source_path.read_bytes())
    output = entry_dir / "preview.xml.txt"
    await asyncio.to_thread(output.write_text, content, encoding="utf-8")
    # Keep the downloaded representation inert. The frontend renders and
    # highlights it as text without inserting untrusted XML into the DOM.
    return ProcessResult("xml", detected.label, "text/plain; charset=utf-8", output.name)


async def _render_json(
    source_path: Path, entry_dir: Path, detected: DetectedFormat, settings: Settings
) -> ProcessResult:
    if source_path.stat().st_size > settings.max_text_bytes:
        raise HTTPException(status_code=413, detail="JSON-файл превышает лимит")
    content = await asyncio.to_thread(_decode_text, source_path.read_bytes())
    output = entry_dir / "preview.json.txt"
    await asyncio.to_thread(output.write_text, content, encoding="utf-8")
    return ProcessResult("json", detected.label, "text/plain; charset=utf-8", output.name)


async def _render_html(
    source_path: Path, entry_dir: Path, detected: DetectedFormat, settings: Settings
) -> ProcessResult:
    if source_path.stat().st_size > settings.max_text_bytes:
        raise HTTPException(status_code=413, detail="HTML-файл превышает лимит")
    content = await asyncio.to_thread(_decode_text, source_path.read_bytes())
    sanitized = _sanitize_email_html(content)
    output = entry_dir / "preview.html"
    await asyncio.to_thread(output.write_text, sanitized, encoding="utf-8")
    return ProcessResult("html", detected.label, "text/html; charset=utf-8", output.name)


async def _convert_image(
    source_path: Path, entry_dir: Path, detected: DetectedFormat
) -> ProcessResult:
    output = entry_dir / "preview.png"

    def convert() -> None:
        try:
            with Image.open(source_path) as image:
                image.seek(0)
                image.thumbnail((12_000, 12_000))
                if image.mode not in {"RGB", "RGBA"}:
                    image = image.convert("RGBA" if "transparency" in image.info else "RGB")
                image.save(output, "PNG", optimize=True)
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
            raise HTTPException(status_code=422, detail="Изображение не удалось прочитать") from exc

    await asyncio.to_thread(convert)
    return ProcessResult("image", detected.label, "image/png", output.name)


async def _run(command: list[str], timeout: int, error_prefix: str) -> None:
    try:
        process = await asyncio.create_subprocess_exec(
            *command,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
    except FileNotFoundError as exc:
        raise HTTPException(status_code=503, detail=f"{error_prefix}: обработчик не установлен") from exc
    try:
        stdout, stderr = await asyncio.wait_for(process.communicate(), timeout=timeout)
    except TimeoutError as exc:
        process.kill()
        await process.wait()
        raise HTTPException(status_code=504, detail=f"{error_prefix}: превышено время") from exc
    if process.returncode != 0:
        details = (stderr or stdout).decode("utf-8", "replace").strip()[-500:]
        raise HTTPException(status_code=422, detail=f"{error_prefix}: {details}")


_DECLARED_ENCODING_PATTERNS = (
    re.compile(br"<\?xml[^>]{0,512}\bencoding\s*=\s*['\"]\s*([a-z0-9._:-]+)", re.IGNORECASE),
    re.compile(br"<meta\b[^>]{0,1024}\bcharset\s*=\s*['\"]?\s*([a-z0-9._:-]+)", re.IGNORECASE),
    re.compile(br"<meta\b[^>]{0,1024}\bcontent\s*=\s*['\"][^'\"]*charset\s*=\s*([a-z0-9._:-]+)", re.IGNORECASE),
)

_ALLOWED_DECLARED_ENCODINGS = {
    "ascii",
    "cp720",
    "cp864",
    "cp1006",
    "cp1251",
    "cp1252",
    "cp1256",
    "iso8859-1",
    "iso8859-6",
    "koi8-r",
    "mac-cyrillic",
    "utf-8",
    "utf-8-sig",
    "utf-16",
    "utf-16-be",
    "utf-16-le",
    "utf-32",
    "utf-32-be",
    "utf-32-le",
}

_LEGACY_ENCODINGS = (
    "cp1256",     # Windows Arabic
    "iso8859-6",  # ISO Arabic
    "cp720",      # DOS Arabic
    "cp1251",     # preserve the existing Cyrillic support
    "cp1252",
    "latin-1",
)


def _decode_text(data: bytes) -> str:
    unicode_encoding = _unicode_encoding(data)
    if unicode_encoding:
        return data.decode(unicode_encoding, "replace")

    declared_encoding = _declared_encoding(data)
    if declared_encoding:
        try:
            return data.decode(declared_encoding)
        except UnicodeDecodeError:
            # A broken declaration should not make the preview unavailable.
            pass

    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass

    candidates: list[tuple[float, int, str]] = []
    for priority, encoding in enumerate(_LEGACY_ENCODINGS):
        try:
            decoded = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        candidates.append((_text_quality(decoded), -priority, decoded))
    if candidates:
        return max(candidates, key=lambda candidate: (candidate[0], candidate[1]))[2]
    return data.decode("utf-8", "replace")


def _unicode_encoding(data: bytes) -> str | None:
    # UTF-32 little-endian starts with the UTF-16 LE BOM, so the longer
    # signatures must be checked first.
    signatures = (
        (b"\x00\x00\xfe\xff", "utf-32"),
        (b"\xff\xfe\x00\x00", "utf-32"),
        (b"\xef\xbb\xbf", "utf-8-sig"),
        (b"\xfe\xff", "utf-16"),
        (b"\xff\xfe", "utf-16"),
        (b"\x00\x00\x00<", "utf-32-be"),
        (b"<\x00\x00\x00", "utf-32-le"),
        (b"\x00<\x00?", "utf-16-be"),
        (b"<\x00?\x00", "utf-16-le"),
    )
    return next((encoding for signature, encoding in signatures if data.startswith(signature)), None)


def _declared_encoding(data: bytes) -> str | None:
    sample = data[:4096]
    for pattern in _DECLARED_ENCODING_PATTERNS:
        match = pattern.search(sample)
        if not match:
            continue
        try:
            requested = match.group(1).decode("ascii")
            canonical = codecs.lookup(requested).name
        except (LookupError, UnicodeDecodeError):
            continue
        if canonical in _ALLOWED_DECLARED_ENCODINGS:
            return canonical
    return None


def _text_quality(value: str) -> float:
    scripts = {"arabic": 0, "cyrillic": 0, "latin": 0}
    controls = 0
    suspicious_symbols = 0
    compatibility_characters = 0
    cyrillic_uppercase = 0
    cyrillic_internal_uppercase = 0
    at_word_start = True

    for character in value:
        category = unicodedata.category(character)
        if category.startswith("C") and character not in "\n\r\t\f":
            controls += 1
        name = unicodedata.name(character, "")
        if name.startswith("ARABIC"):
            scripts["arabic"] += 1
        elif name.startswith("CYRILLIC"):
            scripts["cyrillic"] += 1
            if character.isupper():
                cyrillic_uppercase += 1
                if not at_word_start:
                    cyrillic_internal_uppercase += 1
        elif name.startswith("LATIN"):
            scripts["latin"] += 1
        if name.startswith(("BOX DRAWINGS", "BLOCK ELEMENT", "PRIVATE USE")):
            suspicious_symbols += 1
        if character != unicodedata.normalize("NFKC", character):
            compatibility_characters += 1
        at_word_start = not character.isalpha()

    letter_count = sum(scripts.values())
    dominant = max(scripts.values())
    mixed_scripts = letter_count - dominant
    score = (
        dominant * 5
        - mixed_scripts * 4
        - controls * 30
        - suspicious_symbols * 8
        - compatibility_characters * 8
    )

    # Arabic bytes interpreted as Windows-1251 commonly turn into oddly
    # capitalized Cyrillic words (for example مرحبا -> гСНИЗ). Penalizing
    # capitals inside words separates those encodings without penalizing
    # normal sentence capitalization.
    if scripts["cyrillic"] >= 4:
        score -= cyrillic_internal_uppercase * 8
        if cyrillic_uppercase / scripts["cyrillic"] > 0.65:
            score -= scripts["cyrillic"] * 2
    return score


def _sanitize_email_html(value: str) -> str:
    allowed_tags = {
        "a", "abbr", "b", "blockquote", "br", "caption", "code", "div", "em", "h1", "h2",
        "h3", "h4", "h5", "h6", "hr", "i", "img", "li", "ol", "p", "pre", "small", "span",
        "strong", "sub", "sup", "table", "tbody", "td", "tfoot", "th", "thead", "tr", "u", "ul",
    }

    def allow_attribute(tag: str, name: str, attribute_value: str) -> bool:
        if name in {"colspan", "rowspan", "title", "alt"}:
            return True
        if name == "dir":
            return attribute_value.lower() in {"auto", "ltr", "rtl"}
        if name == "lang":
            return bool(re.fullmatch(r"[A-Za-z]{1,8}(?:-[A-Za-z0-9]{1,8})*", attribute_value))
        if tag == "a" and name == "href":
            return attribute_value.startswith(("http://", "https://", "mailto:"))
        if tag == "img" and name == "src":
            return attribute_value.startswith(("cid:", "data:image/"))
        return False

    return bleach.clean(
        value,
        tags=allowed_tags,
        attributes=allow_attribute,
        protocols={"http", "https", "mailto", "cid", "data"},
        strip=True,
        strip_comments=True,
    )


def zipfile_is_zip(path: Path) -> bool:
    import zipfile

    return zipfile.is_zipfile(path)
